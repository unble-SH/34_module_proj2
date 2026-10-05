#!/usr/bin/env python3
"""
사례 수집 시트(cases.xlsx) → cases.json 변환기

사용법 (report_ai 폴더에서 실행):
    python convert_cases.py
    python convert_cases.py --xlsx data/cases.xlsx --out data/cases.json --sources data/sources

하는 일
  1. cases 시트의 모든 행을 카드(JSON)로 바꾼다. 빈 행은 건너뛴다.
  2. lists 시트를 기준으로 드롭다운 값(type, 팀 취약점, 공격 태그, ISMS-P)을 검사한다.
  3. 필수값, 날짜 형식, case_id 형식·중복, 금액 칸, source_file 존재 여부를 검사한다.
  4. 오류가 하나라도 있으면 JSON을 쓰지 않고 오류 목록만 보여준다. 경고는 보여주고 넘어간다.
  5. law_regime(구법/확인필요)을 날짜로 채우고, 팀 취약점별 커버리지를 출력한다.

note 열은 내부 메모라서 JSON에 넣지 않는다.
verified가 FALSE인 카드도 JSON에는 들어가고, 지식베이스에 적재할 때 걸러낸다.
"""
import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
AMENDED = date(2026, 9, 11)  # 개정 개인정보 보호법 시행일
PREFIX = {"처분": "PIPC", "판례": "COURT", "기사": "NEWS", "심사결함사례": "ISMSP"}
REQUIRED = ["case_id", "type", "date", "org", "attack_tag_1", "outcome", "source_url", "verified"]
HEADERS = ["case_id", "type", "date", "org", "industry", "violated_codes",
           "team_vuln_1", "team_vuln_2", "attack_tag_1", "attack_tag_2", "attack_tag_3",
           "isms_p_1", "isms_p_2", "isms_p_3", "facts", "outcome", "fine_krw", "penalty_krw",
           "source_url", "source_file", "verified", "note"]
OPTIONAL_HEADERS = {"isms_p_3"}  # 예전 시트(2칸)도 읽을 수 있게 없으면 빈 값으로 처리
LIST_COLS = {"type": 1, "team_vuln": 2, "attack_tag": 5, "isms_p": 8}  # lists 시트의 열 번호
CASE_ID_RE = re.compile(r"^(PIPC|COURT|NEWS|ISMSP)-(\d{4})-(\d{3})$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
CODE_RE = re.compile(r"^[A-Z]+-\d+(-\d+)?$")


def nfc(v):
    """맥에서 입력한 한글(NFD)도 같은 글자로 비교되게 NFC로 맞추고, 빈 문자열은 None으로."""
    if isinstance(v, str):
        v = unicodedata.normalize("NFC", v).strip()
        return v or None
    return v


class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    def err(self, row, msg):
        self.errors.append(f"  {row}행: {msg}")

    def warn(self, row, msg):
        self.warnings.append(f"  {row}행: {msg}")


def read_lists(wb):
    ls = wb["lists"]
    out = {}
    for name, col in LIST_COLS.items():
        vals = {nfc(ls.cell(row=r, column=col).value) for r in range(2, ls.max_row + 1)}
        out[name] = {str(v) for v in vals if v is not None}
    return out


def to_date(v, row, rep):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str) and DATE_RE.match(v):
        try:
            datetime.strptime(v, "%Y-%m-%d")
            return v
        except ValueError:
            pass
    rep.err(row, f"date '{v}'는 YYYY-MM-DD 형식의 실제 날짜가 아님")
    return None


def to_int(v, name, row, rep):
    if v is None:
        return None
    if isinstance(v, bool):
        rep.err(row, f"{name}에 TRUE/FALSE가 들어 있음")
        return None
    if isinstance(v, (int, float)) and float(v).is_integer() and v >= 0:
        return int(v)
    if isinstance(v, str) and v.isdigit():
        return int(v)
    rep.err(row, f"{name} '{v}'는 0 이상 정수가 아님 (원·쉼표 없이 숫자만)")
    return None


def to_bool(v, row, rep):
    if isinstance(v, bool):
        return v
    if isinstance(v, str) and v.upper() in ("TRUE", "FALSE"):
        return v.upper() == "TRUE"
    rep.err(row, f"verified '{v}'는 TRUE/FALSE가 아님")
    return None


def convert(xlsx, sources_dir, rep):
    wb = load_workbook(xlsx, data_only=True)
    if "cases" not in wb.sheetnames or "lists" not in wb.sheetnames:
        rep.err(0, "시트 이름은 cases, lists여야 함")
        return []
    ws = wb["cases"]
    allowed = read_lists(wb)

    header = [nfc(c.value) for c in ws[1]]
    missing = [h for h in HEADERS if h not in header and h not in OPTIONAL_HEADERS]
    if missing:
        rep.err(1, f"열 이름이 없음: {', '.join(missing)} (1행 열 이름은 바꾸면 안 됨)")
        return []
    col = {h: header.index(h) for h in HEADERS if h in header}

    source_names = set()
    if sources_dir.is_dir():
        source_names = {unicodedata.normalize("NFC", p.name) for p in sources_dir.iterdir() if p.is_file()}
    else:
        rep.warnings.append(f"  sources 폴더가 없어서 source_file 존재 여부는 건너뜀: {sources_dir}")

    cards, seen = [], {}
    for r, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = {h: None for h in HEADERS}
        d.update({h: nfc(cells[i]) if i < len(cells) else None for h, i in col.items()})
        if all(v is None for v in d.values()):
            continue

        for h in REQUIRED:
            if d[h] is None:
                rep.err(r, f"필수값 {h}가 비어 있음")

        # type과 case_id
        ctype = d["type"]
        if ctype is not None and ctype not in allowed["type"]:
            rep.err(r, f"type '{ctype}'는 lists 시트에 없음")
        cid = d["case_id"]
        if cid is not None:
            m = CASE_ID_RE.match(str(cid))
            if not m:
                rep.err(r, f"case_id '{cid}' 형식이 틀림 (예: PIPC-2025-001)")
            elif ctype in PREFIX and m.group(1) != PREFIX[ctype]:
                rep.err(r, f"case_id 접두어 {m.group(1)}가 type '{ctype}'과 안 맞음 ({PREFIX[ctype]}-여야 함)")
            if cid in seen:
                rep.err(r, f"case_id '{cid}'가 {seen[cid]}행과 중복")
            seen[cid] = r

        cdate = to_date(d["date"], r, rep) if d["date"] is not None else None
        if cid and cdate and CASE_ID_RE.match(str(cid)) and CASE_ID_RE.match(str(cid)).group(2) != cdate[:4]:
            rep.warn(r, f"case_id 연도와 date 연도가 다름 ({cid} / {cdate})")

        # 드롭다운 값
        def pick(names, kind):
            vals = []
            for n in names:
                v = d[n]
                if v is None:
                    continue
                v = str(v)
                if v not in allowed[kind]:
                    rep.err(r, f"{n} '{v}'는 lists 시트에 없음")
                elif v not in vals:
                    vals.append(v)
            return vals

        team_vuln = pick(["team_vuln_1", "team_vuln_2"], "team_vuln")
        attack_tags = pick(["attack_tag_1", "attack_tag_2", "attack_tag_3"], "attack_tag")
        isms_p = pick(["isms_p_1", "isms_p_2", "isms_p_3"], "isms_p")
        if not team_vuln:
            rep.warn(r, "팀 취약점이 비어 있음 → 일반 사례로 처리")

        codes = [c.strip() for c in str(d["violated_codes"]).split(",")] if d["violated_codes"] else []
        for c in codes:
            if not CODE_RE.match(c):
                rep.warn(r, f"violated_codes '{c}'가 코드 형식(예: PIPA-29)이 아님")

        # 금액
        fine = to_int(d["fine_krw"], "fine_krw", r, rep)
        penalty = to_int(d["penalty_krw"], "penalty_krw", r, rep)
        if ctype == "처분":
            if d["fine_krw"] is None:
                rep.err(r, "처분인데 fine_krw가 비어 있음 (과징금이 없으면 0)")
            if d["penalty_krw"] is None:
                rep.err(r, "처분인데 penalty_krw가 비어 있음 (과태료가 없으면 0)")

        # 원문 파일
        sfile = d["source_file"]
        if sfile and source_names and sfile not in source_names:
            rep.err(r, f"source_file '{sfile}'가 sources 폴더에 없음")
        if not sfile and ctype in ("처분", "판례"):
            rep.warn(r, "처분·판례인데 source_file이 비어 있음")
        if sfile and ctype == "기사":
            rep.warn(r, "기사는 원문 파일을 저장하지 않기로 함 (링크·요약만)")

        verified = to_bool(d["verified"], r, rep) if d["verified"] is not None else None

        # 적용 법 체계: 의결일 기준 1차 판별 (시행 후 의결이어도 위반 시점이 시행 전이면 구법)
        regime = None
        if cdate:
            regime = "구법" if date.fromisoformat(cdate) < AMENDED else "확인필요"
            if regime == "확인필요":
                rep.warn(r, "개정법 시행 후 의결 → 위반행위 시점을 원문에서 확인해 구법/개정법 판단 필요")

        cards.append({
            "case_id": cid, "type": ctype, "date": cdate, "org": d["org"], "industry": d["industry"],
            "violated_codes": codes, "team_vuln": team_vuln, "attack_tags": attack_tags,
            "isms_p_ids": isms_p, "facts": d["facts"], "outcome": d["outcome"],
            "fine_krw": fine, "penalty_krw": penalty, "law_regime": regime,
            "source_url": d["source_url"], "source_file": sfile, "verified": verified,
        })
    return cards, allowed


def print_summary(cards, allowed):
    ok = [c for c in cards if c["verified"]]
    print(f"\n카드 {len(cards)}장 (verified TRUE {len(ok)}장 / FALSE {len(cards) - len(ok)}장)")
    print("FALSE 카드는 JSON에는 들어가지만 지식베이스 적재 때 제외돼요.\n")
    cnt_all = Counter(v for c in cards for v in c["team_vuln"])
    cnt_ok = Counter(v for c in ok for v in c["team_vuln"])
    print("팀 취약점별 커버리지 (전체 / verified)")
    for v in sorted(allowed["team_vuln"]):
        flag = "  ← 비어 있음" if cnt_all[v] == 0 else ""
        print(f"  {v:<16} {cnt_all[v]:>2} / {cnt_ok[v]:>2}{flag}")
    general = sum(1 for c in cards if not c["team_vuln"])
    if general:
        print(f"  (팀 취약점 없는 일반 사례 {general}장)")


def main():
    ap = argparse.ArgumentParser(description="사례 수집 시트 → cases.json")
    ap.add_argument("--xlsx", default=BASE / "data" / "cases.xlsx", type=Path)
    ap.add_argument("--out", default=BASE / "data" / "cases.json", type=Path)
    ap.add_argument("--sources", default=BASE / "data" / "sources", type=Path)
    args = ap.parse_args()

    if not args.xlsx.exists():
        sys.exit(f"시트 파일이 없어요: {args.xlsx}")

    rep = Report()
    result = convert(args.xlsx, args.sources, rep)
    cards, allowed = result if isinstance(result, tuple) else (result, None)

    if rep.warnings:
        print(f"경고 {len(rep.warnings)}개 (JSON은 만들어져요)")
        print("\n".join(rep.warnings))
    if rep.errors:
        print(f"\n오류 {len(rep.errors)}개 → JSON을 만들지 않았어요. 시트를 고친 뒤 다시 실행하세요.")
        print("\n".join(rep.errors))
        sys.exit(1)

    cards.sort(key=lambda c: c["case_id"])
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(cards, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n저장 완료: {args.out}")
    print_summary(cards, allowed)


if __name__ == "__main__":
    main()
